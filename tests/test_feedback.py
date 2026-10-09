import unittest
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from docx import Document
import json
from agent.word_export import render_report_docx
from tests.test_word_export import example

class FeedbackTests(unittest.TestCase):
 def test_article_nine_boundary_covers_both_certificate_and_money(self):
  entries=json.loads((Path(__file__).resolve().parents[1]/'knowledge/entries.json').read_text(encoding='utf-8-sig'))
  item=next(e for e in entries if e['id']=='K001')
  self.assertIn('证件',item['boundary']);self.assertIn('核验',item['boundary']);self.assertIn('收费',item['boundary'])
  self.assertIn('flk.npc.gov.cn/detail?',item['url'])
  self.assertIn('id=2c909fdd678bf17901678bf74d7106b3',item['url'])
  self.assertEqual(item['source_publisher'],'全国人民代表大会常务委员会')
  self.assertEqual(item['legal_database_status'],'specific_record_checked')
 def test_report_renders_level_mapping_citations_and_self_checks(self):
  report=example();report['self_checks']=['是否已经实际到岗并提供劳动？']
  document=Document(BytesIO(render_report_docx(report)))
  text='\n'.join(p.text for p in document.paragraphs)
  for value in ['红灯','黄灯','灰色','概览依据：K001','关联自查清单','是否已经实际到岗并提供劳动？']:
   self.assertIn(value.replace(' ',''),text.replace(' ',''))
  self.assertNotIn('未签合同可主张二倍工资',text)
 def test_official_laws_precede_guides_without_changing_findings(self):
  report=example();law=report['sources'][0]
  report['sources']=[dict(law,id='G001',title='证据指南',kind='guide'),law]
  report['findings'][0]['evidence_ids'].append('G001')
  document=Document(BytesIO(render_report_docx(report)))
  text='\n'.join(p.text for p in document.paragraphs)
  self.assertLess(text.index('K001  中华人民共和国劳动合同法'),text.index('G001  证据指南'))
 def test_employment_agreement_prompts_checks_without_adding_risks(self):
  from agent.feedback import prepare_report
  report=example();before=list(report['findings'])
  prepare_report(report,'毕业生签订就业协议，学校扣留毕业证原件。')
  self.assertTrue(any('到岗' in q for q in report['self_checks']))
  self.assertEqual(report['findings'],before)
 def test_known_preemployment_status_is_not_asked_again(self):
  from agent.feedback import prepare_report
  report=example();prepare_report(report,'签订就业协议，但尚未入职、尚未到岗。')
  self.assertFalse(any('是否已经实际到岗' in q for q in report['self_checks']))
 def test_unprovided_employment_agreement_is_a_check(self):
  from agent.report_layout import normalize_layout
  from agent.schema import TurnResult,Report
  data=example();data['findings'][0].update(title='就业协议与入职安排条款不明，缺少书面用工约定',level='uncertain')
  result=TurnResult(action='report',message='',questions=[],evidence_ids=['K001'],report=Report(**{k:v for k,v in data.items() if k in Report.model_fields}))
  normalize_layout(result,{'current_message':'签订就业协议，约定入职从事教学工作。','history':[],'materials':[]})
  self.assertEqual(result.report.findings,[]);self.assertTrue(result.report.limitations)

if __name__=='__main__':unittest.main()
