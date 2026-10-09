import unittest
from io import BytesIO
from docx import Document
from agent.word_export import render_report_docx
from tests.test_word_export import example
class OverviewTableTests(unittest.TestCase):
    def test_table_lists_high_and_attention_only_after_priority(self):
        report=example();base=report['findings'][0]
        report['findings']=[{**base,'level':'high','title':'收费培训','suggestions':['先暂停付款，书面确认费用安排。']},{**base,'level':'attention','title':'岗位职责不清','suggestions':['要求提供岗位职责清单。']},{**base,'level':'uncertain','title':'贷款性质待核实'}]
        doc=Document(BytesIO(render_report_docx(report)))
        self.assertEqual(len(doc.tables),1)
        rows=[[cell.text for cell in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(rows,[['风险等级','问题','建议先做'],['高风险','收费培训','先暂停付款，书面确认费用安排。'],['需关注','岗位职责不清','要求提供岗位职责清单。']])
        body=list(doc.element.body);table=doc.tables[0]._tbl
        priority=next(p._p for p in doc.paragraphs if p.text.startswith('当前优先行动'))
        detail=next(p._p for p in doc.paragraphs if p.text=='02  风险详情')
        self.assertLess(body.index(priority),body.index(table));self.assertLess(body.index(table),body.index(detail))
    def test_no_empty_table_when_only_pending_items(self):
        report=example();report['findings'][0]['level']='uncertain'
        doc=Document(BytesIO(render_report_docx(report)))
        self.assertEqual(len(doc.tables),0)
