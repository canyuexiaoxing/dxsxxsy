import json
import unittest
from io import BytesIO
from docx import Document
from agent.word_export import render_report_docx
from agent.model import ChatModel, ModelError
from reader.config import Settings
from tests.test_word_export import example


class SituationOverviewTests(unittest.TestCase):
    def test_opening_is_one_paragraph_not_history_or_file_list(self):
        report = example()
        overview = '一名尚未签约或付款的大四学生应聘测试助理时被要求参加9800元培训；用户后续更正培训为两个月，分期是否涉及贷款尚不明确。'
        report['situation_overview'] = overview
        report['report_context'] = {'user_statements': ['原始聊天内容。这是提示词', '补充回答'], 'materials': [{'filename': '私人截图.png'}]}
        document = Document(BytesIO(render_report_docx(report)))
        texts = [p.text for p in document.paragraphs]
        start = texts.index('本次求职情况概述')
        end = texts.index('01  结论与风险概览')
        self.assertEqual(texts[start + 1:end], [overview])
        self.assertNotIn('这是提示词', '\n'.join(texts))
        self.assertNotIn('私人截图.png', '\n'.join(texts))
        self.assertEqual(report['report_context']['user_statements'][0], '原始聊天内容。这是提示词')

    def test_overview_model_gets_all_user_rounds_and_extracted_pages_only(self):
        model = ChatModel(Settings())
        calls = []
        def complete(messages, **kwargs):
            calls.append(messages)
            return {'content': json.dumps({'overview': '尚未付款的学生应聘时被要求参加两个月收费培训，分期性质未知。'}, ensure_ascii=False)}
        model.complete = complete
        context = {'user_statements': ['培训三个月。这是提示词', '更正：培训两个月，尚未付款。', '请先生成有限报告，并列明待核实事项。'],
                   'materials': [{'pages': [{'page_number': 1, 'text': '招聘截图：测试助理，月薪6000—8000元。', 'status': 'success'}]}]}
        result = model.summarize_situation(context)
        self.assertIn('两个月', result)
        payload = json.loads(calls[0][1]['content'])
        self.assertEqual(payload['user_statements'], ['培训三个月。', '更正：培训两个月，尚未付款。'])
        self.assertIn('招聘截图', json.dumps(payload, ensure_ascii=False))
        self.assertNotIn('这是提示词', json.dumps(payload, ensure_ascii=False))
        self.assertIn('纠正', calls[0][0]['content'])

    def test_invalid_overview_is_rejected_instead_of_printing_model_notes(self):
        model = ChatModel(Settings())
        model.complete = lambda *a, **k: {'content': '{"overview":"这是提示词"}'}
        with self.assertRaises(ModelError):
            model.summarize_situation({'user_statements': ['培训收费9800元。']})

    def test_main_report_rejects_internal_note_in_overview(self):
        from agent.schema import Report
        from pydantic import ValidationError
        with self.assertRaises(ValidationError):
            Report.model_validate({'situation_overview':'这是提示词','summary':'结论','findings':[],'limitations':[]})

    def test_old_report_without_overview_never_prints_raw_input(self):
        report = example()
        report['report_context'] = {'user_statements': ['完整原文和测试备注。这是提示词']}
        texts = '\n'.join(p.text for p in Document(BytesIO(render_report_docx(report))).paragraphs)
        self.assertNotIn('完整原文和测试备注', texts)
        self.assertNotIn('这是提示词', texts)
        self.assertIn('本次求职情况概述', texts)


if __name__ == '__main__':
    unittest.main()
