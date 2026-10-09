import sys, unittest
from pathlib import Path
from io import BytesIO
from docx import Document
from agent.word_export import render_report_docx
from tests.test_word_export import example

class FourthFeedbackTests(unittest.TestCase):
    def texts(self,r):
        d=Document(BytesIO(render_report_docx(r)))
        return '\n'.join(p.text for p in d.paragraphs)
    def test_counts_distinguish_limits_from_risks(self):
        r=example();r['limitations']=['主体未知','协议未见']
        self.assertIn('第03节的2条待核实事项与分析局限',self.texts(r))
    def test_pending_card_counts_section03_instead_of_uncertain_findings(self):
        from unittest.mock import patch
        from PIL import ImageDraw
        original_draw=ImageDraw.Draw
        rendered=[]
        def capture(image):
            draw=original_draw(image)
            original_text=draw.text
            def text(position,value,*args,**kwargs):
                rendered.append(str(value))
                return original_text(position,value,*args,**kwargs)
            draw.text=text
            return draw
        r=example();r['limitations']=[f'待核实事项{i}' for i in range(7)]
        with patch('agent.word_export.ImageDraw.Draw',side_effect=capture):
            text=self.texts(r)
        self.assertEqual([value for value in rendered if value.isdigit()],['1','0','7'])
        self.assertIn('待核实事项与局限',rendered)
        self.assertIn('第03节条目数',rendered)
        self.assertIn('待核实卡片统计第03节的7条待核实事项与分析局限',text)

    def test_overview_order_keeps_actions_and_reading_note_before_cards(self):
        r=example();r['summary']='押金和证件保管要求需要优先处理。应在接受岗位前核实相关条件。\n\n薪资和实际岗位不一致，也会影响求职决定。需要取得书面说明。'
        r['prevention']={'label':'暂缓付款、签约和交证件','steps':['不要支付押金或交证件原件。','书面确认薪资和实际岗位。'],'note':'不构成违法认定。'}
        d=Document(BytesIO(render_report_docx(r)));texts=[p.text for p in d.paragraphs]
        start=texts.index('01  结论与风险概览');end=texts.index('02  风险详情')
        opening=texts[start:end]
        self.assertTrue(opening[1].startswith('整体建议'))
        priority=next(i for i,t in enumerate(opening) if t.startswith('当前优先行动'))
        reading=next(i for i,t in enumerate(opening) if '各项原因、依据和具体建议见第02节' in t)
        drawing=next(i for i,p in enumerate(d.paragraphs[start:end]) if p._p.xpath('.//w:drawing'))
        self.assertLess(priority,reading);self.assertLess(reading,drawing)
        self.assertIn('问题速览',texts)
        self.assertNotIn('签约与入职前，建议先做',opening)
        self.assertNotIn('书面确认薪资和实际岗位。',opening)
        self.assertIn('书面确认薪资和实际岗位。',texts[end:])

    def test_unused_sources_do_not_get_cards(self):
        r=example();r['sources'].append({**r['sources'][0],'id':'unused','title':'未引用卡片','text':'不应展示的案件处理结果'})
        self.assertNotIn('不应展示的案件处理结果',self.texts(r))
    def test_testing_note_separate_and_original_not_mutated(self):
        r=example();r['report_context']={'user_statements':['培训费9800元，可以分期。这是提示词']}
        text=self.texts(r)
        self.assertIn('本次求职情况概述',text)
        self.assertNotIn('培训费9800元，可以分期。',text)
        self.assertNotIn('分析请求与测试说明',text)
        self.assertNotIn('这是提示词',text)
        self.assertEqual(r['report_context']['user_statements'][0],'培训费9800元，可以分期。这是提示词')
    def test_repeated_source_details_point_to_named_source(self):
        r=example();r['sources'][0].update(source_publisher='全国人大常委会',source_site='国家法律法规数据库',source_access_note='原始记录已核对')
        r['sources'].append({**r['sources'][0],'id':'K002','title':'另一部法律','url':'https://example.gov.cn/another'})
        r['findings'][0]['evidence_ids'].append('K002')
        text=self.texts(r)
        self.assertNotIn('05  来源核对详情',text)
        self.assertEqual(text.count('全国人大常委会'),1)
        self.assertEqual(text.count('原始记录已核对'),1)
        self.assertIn('同 K001',text)

    def test_risk_number_and_source_detail_appendix(self):
        r=example();r['sources'][0].update(source_site='国家法律法规数据库',source_publisher='全国人大常委会')
        text=self.texts(r)
        self.assertIn('风险1',text)
        self.assertNotIn('05  来源核对详情',text)
        self.assertIn('全国人大常委会',text)
        self.assertIn('国家法律法规数据库',text)
        self.assertNotIn('https://example.gov.cn/law',text)

if __name__=='__main__':unittest.main()
