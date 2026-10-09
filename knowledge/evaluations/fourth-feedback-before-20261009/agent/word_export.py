"""Render structured reports as editable Word documents; no model calls."""
from .feedback import LEVEL_NOTE, ordered_sources, overview_ids
from .presentation import public_report, URL_PATTERN
import re
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit, quote

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

LEVELS={'high':('高风险','B85646'),'attention':('需关注','B78224'),'uncertain':('待核实','69798B')}
GREEN='287B65'

def risk_chart(findings):
    counts=Counter(f.get('level') for f in findings)
    image=Image.new('RGB',(1500,300),'white'); draw=ImageDraw.Draw(image)
    font_path=Path('C:/Windows/Fonts/msyh.ttc')
    def font(size):
        return ImageFont.truetype(str(font_path),size) if font_path.exists() else ImageFont.load_default(size=size)
    for i,(level,(label,color)) in enumerate(LEVELS.items()):
        x=10+i*500
        draw.rounded_rectangle((x,10,x+475,286),radius=24,fill='#F5F8F7')
        draw.rounded_rectangle((x+24,35,x+35,247),radius=5,fill='#'+color)
        draw.text((x+65,38),label,font=font(32),fill='#'+color)
        draw.text((x+65,91),str(counts[level]),font=font(76),fill='#1D3430')
        draw.text((x+65,208),'项已标注的问题',font=font(25),fill='#6A7C77')
    buffer=BytesIO(); image.save(buffer,format='PNG'); buffer.seek(0)
    return buffer

def shade(paragraph,color):
    element=OxmlElement('w:shd'); element.set(qn('w:fill'),color)
    paragraph._p.get_or_add_pPr().append(element)

def link(paragraph,text,url):
    if urlsplit(url).scheme not in ('http','https'): return
    url=quote(url,safe=":/?&=#%+@,;~-._!$'()*[]")
    relationship=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    element=OxmlElement('w:hyperlink'); element.set(qn('r:id'),relationship)
    element.set(qn('w:history'),'1')
    run=OxmlElement('w:r'); properties=OxmlElement('w:rPr')
    style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');properties.append(style)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
    color=OxmlElement('w:color'); color.set(qn('w:val'),GREEN); properties.append(color)
    run.append(properties); value=OxmlElement('w:t'); value.text=text; run.append(value); element.append(run)
    paragraph._p.append(element)

def render_report_docx(report,title='求职材料自查',material_names=None):
    report=public_report(report)
    names=material_names or {}; document=Document(); section=document.sections[0]
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.top_margin=Cm(1.9); section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(2); section.right_margin=Cm(2)
    section.header_distance=Cm(.8); section.footer_distance=Cm(.8)
    for name in ('Normal','Title','Heading 1','Heading 2','List Bullet'):
        style=document.styles[name]; style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=document.styles['Normal']; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string('263D36')
    normal.paragraph_format.line_spacing=1.4; normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.space_before=Pt(0)
    normal.paragraph_format.keep_together=False
    document.styles['Title'].font.color.rgb=RGBColor.from_string('1D3430')
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Heading 1',17),('Heading 2',12.5)]:
        style=document.styles[name]; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(GREEN)
        style.paragraph_format.space_before=Pt(15); style.paragraph_format.space_after=Pt(8)
    header=section.header.paragraphs[0]; header.text='求职避坑镜  /  JOB MIRROR'
    header.runs[0].font.size=Pt(9); header.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('求职风险自查报告  ·  第 ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    for run in footer.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('71847D')
    document.core_properties.title='求职风险自查报告'; document.core_properties.author='求职避坑镜'

    def paragraph(text='',label=None,color=None):
        p=document.add_paragraph()
        if label: p.add_run(label+'  ').bold=True
        for i,line in enumerate(str(text).splitlines() or ['']):
            if i: p.add_run().add_break()
            run=p.add_run(line)
            if color: run.font.color.rgb=RGBColor.from_string(color)
        return p

    def advice(text,p=None):
        if p is None:p=document.add_paragraph(style='List Bullet')
        end=0
        for match in re.finditer(URL_PATTERN,text):
            p.add_run(text[end:match.start()]);link(p,match.group(),match.group());end=match.end()
        p.add_run(text[end:])
        return p

    document.add_heading('求职风险自查报告',0)
    paragraph('根据已提供的招聘信息、沟通内容与拟签材料进行风险自查。')
    date=report.get('created_at','')
    try: date=datetime.fromisoformat(date.replace('Z','+00:00')).astimezone().strftime('%Y-%m-%d %H:%M')
    except (ValueError,AttributeError): pass
    paragraph(f'第 {report.get("version",1)} 版  |  生成时间：{date or "未记录"}',color='71847D')
    paragraph('本报告依据已提供材料进行风险提示；风险等级不是违法认定或风险概率。',color='71847D')
    context=report.get('report_context',{})
    document.add_heading('分析材料与用户说明',1)
    statements=context.get('user_statements',[])
    if statements:
        paragraph('以下保留生成本版报告前用户提供的完整文字，包括材料描述与分析请求；这是用户陈述，尚未外部核验。后续补充或纠正以较新说明为准。',color='71847D')
        for i,text in enumerate(statements,1):
            paragraph(f'第{i}次用户说明',color='71847D').paragraph_format.keep_with_next=True
            paragraph(text)
    else:
        paragraph('未保存本版完整用户说明；请结合下方材料原文与原始文件理解判断范围。不能将会话短标题视为事实摘要。',color='71847D')
    for material in context.get('materials',[]):
        paragraph(material.get('filename') or material.get('material_id','上传材料'),label='上传材料')
    document.add_heading('01  结论与风险概览',1)
    paragraph(report.get('summary','暂无结论'))
    if report.get('prevention'):
        prevention=report['prevention']
        document.add_heading('签约与入职前，建议先做',2)
        paragraph(prevention['label'])
        for step in prevention['steps']:advice(step)
        paragraph(prevention['note'],color='71847D')
    ids=overview_ids(report)
    if ids: paragraph('、'.join(ids),label='概览依据：',color='71847D')
    paragraph(LEVEL_NOTE,color='71847D')
    findings=report.get('findings',[])
    picture=document.add_paragraph(); picture.add_run().add_picture(risk_chart(findings),width=Cm(17))
    paragraph('图中数量来自本报告已标注风险点；未发现问题不代表已证明安全。',color='71847D')
    if findings:
        document.add_heading('问题速览',2)
        table=document.add_table(rows=1, cols=3); table.autofit=False
        borders=OxmlElement('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        table._tbl.tblPr.append(borders)
        widths=[Cm(2.2),Cm(5.2),Cm(9.6)]
        for column,width in zip(table.columns,widths): column.width=width
        for cell,label,width in zip(table.rows[0].cells,['风险等级','问题','建议先做'],widths):
            cell.width=width; cell.text=label
            element=OxmlElement('w:shd'); element.set(qn('w:fill'),'EDF6F1'); cell._tc.get_or_add_tcPr().append(element)
        repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for finding in findings[:6]:
            row=table.add_row(); level,color=LEVELS.get(finding.get('level'),LEVELS['uncertain'])
            values=[level,finding['title'],next(iter(finding.get('suggestions',[])),'详见风险详情')]
            for index,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
                cell.width=width
                if index==2:advice(value,cell.paragraphs[0])
                else:cell.text=value
            row.cells[0].paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(color)
            no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.2; p.paragraph_format.space_after=Pt(7); p.paragraph_format.space_before=Pt(7)
                    for run in p.runs: run.font.size=Pt(9.5)
        if len(findings)>6: paragraph('速览展示前6项，完整内容见风险详情。',color='71847D')
    if report.get('version',1)>1 and report.get('change_reason'): paragraph(report['change_reason'],label='本版修改')
    document.add_heading('02  风险详情',1)
    sources={s['id']:s for s in report.get('sources',[])}
    if not findings: paragraph('本次未列出具体风险点。请结合分析局限理解结论。')
    for i,finding in enumerate(findings,1):
        start=len(document.paragraphs)
        level,color=LEVELS.get(finding.get('level'),LEVELS['uncertain'])
        p=document.add_heading(f'{i:02d}  {finding["title"]}  ·  {level}',2)
        p.runs[0].font.color.rgb=RGBColor.from_string(color)
        p=paragraph(finding['quote'],label='材料原文'); shade(p,'F3F6F5')
        mid=finding.get('material_id',''); page=finding.get('page_number',0)
        location='用户补充说明' if mid=='user' else names.get(mid,mid)+(f' · 第{page}页' if page else '')
        paragraph(location,label='原文位置',color='71847D')
        paragraph(finding['explanation'],label='风险说明')
        for sid in finding.get('evidence_ids',[]):
            source=sources.get(sid)
            paragraph(f'{sid} · {source["title"]} · {source["location"]}' if source else sid,label='分析依据')
        paragraph('建议行动',color=GREEN).paragraph_format.keep_with_next=True
        for suggestion in finding.get('suggestions',[]): advice(suggestion)
        group=document.paragraphs[start:]
        # Keep headings with their first paragraph, but allow long risk details to span pages.
        p=document.paragraphs[start]
        p.paragraph_format.keep_with_next=True
    document.add_heading('03  待核实事项与分析局限',1)
    for item in report.get('limitations',[]): document.add_paragraph(item,style='List Bullet')
    if not report.get('limitations'): paragraph('原报告未单独列出局限；这不代表所有事实均已核实。')
    if report.get('self_checks'):
        document.add_heading('关联自查清单',2)
        paragraph('可在当前会话补充这些信息，必要时再修订报告；未确认的情形不计为新增风险。',color='71847D')
        for question in report['self_checks']: document.add_paragraph(question,style='List Bullet')
    document.add_heading('官方核查入口',2)
    paragraph('以下为查询入口，报告没有代你查询企业或核验资质。企业登记正常也不能证明岗位安全。',color='71847D')
    p=paragraph(label='企业登记查询');link(p,'国家企业信用信息公示系统 https://www.gsxt.gov.cn/', 'https://www.gsxt.gov.cn/')
    p=paragraph(label='人社咨询与办事');link(p,'全国人力资源和社会保障政务服务平台 https://www.12333.gov.cn/', 'https://www.12333.gov.cn/')
    paragraph('人社咨询电话：12333。具体业务以所在地人社部门受理范围为准。',color='71847D')
    document.add_heading('04  依据与来源',1)
    paragraph('法条、风险提示、指南及案例分别标注；案例结论只用于个案参考。',color='71847D')
    kinds={'law':'法律条文','warning':'官方风险提示','guide':'办事与证据指南','case':'案例参考','support':'辅助资料（团队整理）'}
    linked_urls=set();groups={}
    for source in ordered_sources(report):
        key=(source.get('url') or source['id'],source.get('kind'),source.get('source_publisher'),source.get('source_relation'))
        groups.setdefault(key,[]).append(source)
    for entries in groups.values():
        start=len(document.paragraphs);source=entries[0]
        document.add_heading(f'{source["id"]}  {source["title"]}',2)
        paragraph(kinds.get(source.get('kind'),'资料'),color='71847D')
        if source.get('source_publisher'):paragraph(source['source_publisher'],label=source.get('source_role','发布机构'),color='71847D')
        if source.get('source_site'):paragraph(source['source_site'],label='原文载体',color='71847D')
        dates=sorted({e['verified_at'] for e in entries if e.get('verified_at')})
        if dates:paragraph('、'.join(dates),label='资料核对日期',color='71847D')
        if source.get('source_access_note'):paragraph(source['source_access_note'],label='访问说明',color='71847D')
        if source.get('url'):
            url=source['url']
            if url not in linked_urls:
                label='关联官方依据' if source.get('source_relation')=='related_official_basis' else '原文链接'
                p=paragraph(label=label);link(p,'点击打开官方原文',url)
                p=paragraph();link(p,url,url);linked_urls.add(url)
                for alternate in source.get('alternate_urls',[]):
                    p=paragraph(label='备用官方页面');link(p,alternate,alternate)
            else:paragraph('关联官方原文见前述同一网址的原文链接。',color='71847D')
        elif source.get('source_document'):paragraph(source['source_document']+' · '+source.get('source_paragraphs',''),label='本地来源')
        if source.get('legal_database_url') and source.get('legal_database_status')!='specific_record_checked':
            p=paragraph(label='法律数据库查询入口');link(p,'国家法律法规数据库（查询入口，未核验具体记录）',source['legal_database_url'])
        boundaries={}
        for entry in entries:
            name=entry['title']+' · ' if entry['title']!=source['title'] else ''
            p=paragraph(f'{entry["id"]} · {name}{entry["location"]}')
            for run in p.runs:run.bold=True
            p.paragraph_format.keep_with_next=True
            if entry.get('text'):paragraph(entry['text'])
            if entry.get('boundary'):boundaries.setdefault(entry['boundary'],[]).append(entry['id'])
            if entry.get('verification_status')=='local_document_checked_official_pending':paragraph('仅对照用户提供的本地文书，官方真实性与完整性待核验。',label='核验状态',color='B78224')
        for boundary,ids in boundaries.items():paragraph(boundary,label='适用边界 '+ '、'.join(ids))
        for p in document.paragraphs[start+1:]:
            p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.size=Pt(9.5)
        document.paragraphs[start].paragraph_format.keep_with_next=True
    buffer=BytesIO(); document.save(buffer); return buffer.getvalue()
