"""Render structured reports as editable Word documents; no model calls."""
from .feedback import LEVEL_NOTE, ordered_sources, overview_ids
from .presentation import public_report, URL_PATTERN
import re
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit, quote

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

LEVELS={'high':('高风险','B85646'),'attention':('需关注','B78224'),'uncertain':('待核实','69798B')}
GREEN='287B65'

def risk_chart(findings,limitations):
    counts=Counter(f.get('level') for f in findings)
    image=Image.new('RGB',(1500,140),'white'); draw=ImageDraw.Draw(image)
    font_path=Path('C:/Windows/Fonts/msyh.ttc')
    def font(size):
        return ImageFont.truetype(str(font_path),size) if font_path.exists() else ImageFont.load_default(size=size)
    for i,(level,(label,color)) in enumerate(LEVELS.items()):
        x=10+i*500
        draw.rounded_rectangle((x,10,x+475,130),radius=24,fill='#F5F8F7')
        draw.rounded_rectangle((x+24,25,x+35,115),radius=5,fill='#'+color)
        draw.text((x+65,25),label if level!='uncertain' else '待核实事项与局限',font=font(28),fill='#'+color)
        draw.text((x+340,40),str(len(limitations) if level=='uncertain' else counts[level]),font=font(52),fill='#1D3430')
        draw.text((x+65,75),'第03节条目数' if level=='uncertain' else '已标注风险点',font=font(23),fill='#6A7C77')
    buffer=BytesIO(); image.save(buffer,format='PNG'); buffer.seek(0)
    return buffer

def shade(paragraph,color):
    element=OxmlElement('w:shd'); element.set(qn('w:fill'),color)
    paragraph._p.get_or_add_pPr().append(element)

def link(paragraph,text,url):
    if urlsplit(url).scheme not in ('http','https'): return
    url=quote(url,safe=":/?&=#%+@,;~-._!$'()*[]")
    relationship=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    element=OxmlElement('w:hyperlink'); element.set(qn('r:id'),relationship)
    element.set(qn('w:history'),'1')
    run=OxmlElement('w:r'); properties=OxmlElement('w:rPr')
    style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');properties.append(style)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
    color=OxmlElement('w:color'); color.set(qn('w:val'),GREEN); properties.append(color)
    run.append(properties); value=OxmlElement('w:t'); value.text=text; run.append(value); element.append(run)
    paragraph._p.append(element)

def render_report_docx(report,title='求职材料自查',material_names=None):
    report=public_report(report)
    names=material_names or {}; document=Document(); section=document.sections[0]
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.top_margin=Cm(1.9); section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(2); section.right_margin=Cm(2)
    section.header_distance=Cm(.8); section.footer_distance=Cm(.8)
    for name in ('Normal','Title','Heading 1','Heading 2','List Bullet'):
        style=document.styles[name]; style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=document.styles['Normal']; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string('263D36')
    normal.paragraph_format.line_spacing=1.4; normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.space_before=Pt(0)
    normal.paragraph_format.keep_together=False
    document.styles['Title'].font.color.rgb=RGBColor.from_string('1D3430')
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Heading 1',17),('Heading 2',12.5)]:
        style=document.styles[name]; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(GREEN)
        style.paragraph_format.space_before=Pt(15); style.paragraph_format.space_after=Pt(8)
    header=section.header.paragraphs[0]; header.text='求职避坑镜  /  JOB MIRROR'
    header.runs[0].font.size=Pt(9); header.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('求职风险自查报告  ·  第 ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    for run in footer.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('71847D')
    document.core_properties.title='求职风险自查报告'; document.core_properties.author='求职避坑镜'

    def paragraph(text='',label=None,color=None):
        p=document.add_paragraph()
        if label: p.add_run(label+'  ').bold=True
        for i,line in enumerate(str(text).splitlines() or ['']):
            if i: p.add_run().add_break()
            run=p.add_run(line)
            if color: run.font.color.rgb=RGBColor.from_string(color)
        return p

    def advice(text,p=None):
        if p is None:p=document.add_paragraph(style='List Bullet')
        # Centralize only the two entrances explicitly printed in section 03.
        if 'www.gsxt.gov.cn' in text or 'www.12333.gov.cn' in text:
            text=re.sub(r'\s*核查入口：https://www\.gsxt\.gov\.cn/（国家企业信用信息公示系统）。','',text)
            text=re.sub(r'\s*人社办事入口：https://www\.12333\.gov\.cn/；咨询热线：12333。可选择所在地，具体许可查询和受理渠道请向属地人社部门确认。','',text)
            text+='（核查地址及12333热线见03节“官方核查入口”。）'
        end=0
        for match in re.finditer(URL_PATTERN,text):
            p.add_run(text[end:match.start()]);link(p,match.group(),match.group());end=match.end()
        p.add_run(text[end:])
        return p

    document.add_heading('求职风险自查报告',0)
    paragraph('根据已提供的招聘信息、沟通内容与拟签材料进行风险自查。')
    date=report.get('created_at','')
    try: date=datetime.fromisoformat(date.replace('Z','+00:00')).astimezone().strftime('%Y-%m-%d %H:%M')
    except (ValueError,AttributeError): pass
    paragraph(f'第 {report.get("version",1)} 版  |  生成时间：{date or "未记录"}',color='71847D')
    paragraph('本报告依据已提供材料进行风险提示；风险等级不是违法认定或风险概率。',color='71847D')
    document.add_heading('本次求职情况概述',1)
    from .situation import valid_overview
    overview=report.get('situation_overview','')
    paragraph(' '.join(overview.split()) if valid_overview(overview) else
              '本报告针对已提供的求职材料进行分析；具体要求及尚待核实的信息见下文。')
    document.add_heading('01  结论与风险概览',1)
    findings=report.get('findings',[])
    prevention=report.get('prevention') or {}
    p=paragraph(prevention.get('label') or '先核实材料中的关键条件，再决定是否继续推进。',label='整体建议')
    for run in p.runs:run.bold=True
    summary=str(report.get('summary','暂无结论')).strip()
    blocks=[block.strip() for block in re.split(r'\n\s*\n',summary) if block.strip()]
    if len(blocks)==1:
        sentences=re.findall(r'[^。！？\n]+[。！？]?',summary)
        blocks=[''.join(sentences[i:i+2]).strip() for i in range(0,len(sentences),2)]
    for block in blocks:paragraph(block)
    steps=prevention.get('steps',[])
    priority=next(iter(steps),next((advice for f in findings for advice in f.get('suggestions',[])),
                   '补充影响当前判断的关键信息，再决定是否继续推进。'))
    advice(priority,paragraph(label='当前优先行动'))
    ids=overview_ids(report)
    if ids:paragraph('、'.join(ids),label='概览依据：',color='71847D')
    paragraph('各项原因、依据和具体建议见第02节；尚不明确的信息见第03节。风险等级属于风险提示，不等于违法认定，未检出风险不代表已证明安全。',color='71847D')
    paragraph(LEVEL_NOTE,color='71847D')
    paragraph(f'高风险与需关注卡片统计已标注风险点；待核实卡片统计第03节的{len(report.get("limitations",[]))}条待核实事项与分析局限。',color='71847D')
    picture=document.add_paragraph(); picture.add_run().add_picture(risk_chart(findings,report.get("limitations",[])),width=Cm(17))
    overview_findings=[f for f in findings if f.get('level') in {'high','attention'}]
    if overview_findings:
        document.add_heading('问题速览',2)
        table=document.add_table(rows=1,cols=3);table.autofit=False
        widths=[Cm(2.2),Cm(5.2),Cm(9.6)]
        for column,width in zip(table.columns,widths):column.width=width
        for cell,text in zip(table.rows[0].cells,['风险等级','问题','建议先做']):cell.text=text
        header_repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
        for finding in overview_findings:
            cells=table.add_row().cells
            cells[0].text=LEVELS[finding['level']][0]
            cells[1].text=finding['title']
            advice(next(iter(finding.get('suggestions',[])),'具体行动见第02节对应风险。'),cells[2].paragraphs[0])
            for run in cells[0].paragraphs[0].runs:run.font.color.rgb=RGBColor.from_string(LEVELS[finding['level']][1])
        for index,row in enumerate(table.rows):
            for cell,width in zip(row.cells,widths):
                cell.width=width
                props=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D8E2DC');borders.append(border)
                props.append(borders)
                if index==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'ECF6F1');props.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.35;p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
                    for run in p.runs:
                        run.font.size=Pt(10);run.bold=index==0
    if report.get('version',1)>1 and report.get('change_reason'): paragraph(report['change_reason'],label='本版修改')
    document.add_heading('02  风险详情',1)
    sources={s['id']:s for s in report.get('sources',[])}
    if not findings: paragraph('本次未列出具体风险点。请结合分析局限理解结论。')
    for i,finding in enumerate(findings,1):
        start=len(document.paragraphs)
        level,color=LEVELS.get(finding.get('level'),LEVELS['uncertain'])
        p=document.add_heading(f'风险{i}  {finding["title"]}',2)
        p.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
        paragraph(level,label='风险等级',color=color).paragraph_format.keep_with_next=True
        p=paragraph(finding['quote'],label='材料原文'); shade(p,'F3F6F5')
        mid=finding.get('material_id',''); page=finding.get('page_number',0)
        location='用户补充说明' if mid=='user' else names.get(mid,mid)+(f' · 第{page}页' if page else '')
        paragraph(location,label='原文位置',color='71847D')
        paragraph(finding['explanation'],label='风险说明')
        evidence=[]
        for sid in finding.get('evidence_ids',[]):
            source=sources.get(sid)
            evidence.append(f'{sid} · {source["title"]} · {source["location"]}' if source else sid)
        if evidence:
            p=paragraph(label='分析依据')
            for j,sid in enumerate(finding.get('evidence_ids',[])):
                if j:p.add_run('；')
                element=OxmlElement('w:hyperlink');element.set(qn('w:anchor'),'source_'+sid.replace('-','_'))
                run=OxmlElement('w:r');value=OxmlElement('w:t');value.text=evidence[j];run.append(value);element.append(run);p._p.append(element)
        paragraph('建议行动',color=GREEN).paragraph_format.keep_with_next=True
        groups={'立即采取':[],'沟通与核查':[],'出现以下情况时':[]}
        for suggestion in finding.get('suggestions',[]):
            category='出现以下情况时' if re.match(r'如|若|如果|退出信号',suggestion) else ('立即采取' if re.match(r'今天|立即|先暂停|暂停|不要|不支付|保留|保存',suggestion) else '沟通与核查')
            groups[category].append(suggestion)
        for label,items in groups.items():
            if not items:continue
            paragraph(label,color='71847D').paragraph_format.keep_with_next=True
            for suggestion in items:advice(suggestion)
        group=document.paragraphs[start:]
        # Keep headings with their first paragraph, but allow long risk details to span pages.
        p=document.paragraphs[start]
        p.paragraph_format.keep_with_next=True
    existing_actions={str(item).strip() for finding in findings for item in finding.get('suggestions',[])}
    extra_actions=[]
    for step in steps[1:]:
        if str(step).strip() not in existing_actions:
            extra_actions.append(step);existing_actions.add(str(step).strip())
    if extra_actions:
        document.add_heading('补充行动建议',2)
        for step in extra_actions:advice(step)
    document.add_heading('03  待核实事项与分析局限',1)
    for item in report.get('limitations',[]): document.add_paragraph(item,style='List Bullet')
    if not report.get('limitations'): paragraph('原报告未单独列出局限；这不代表所有事实均已核实。')
    if report.get('self_checks'):
        document.add_heading('关联自查清单',2)
        paragraph('可在当前会话补充这些信息，必要时再修订报告；未确认的情形不计为新增风险。',color='71847D')
        for question in report['self_checks']: document.add_paragraph(question,style='List Bullet')
    document.add_heading('官方核查入口',2)
    paragraph('以下为查询入口，报告没有代你查询企业或核验资质。企业登记正常也不能证明岗位安全。',color='71847D')
    p=paragraph(label='企业登记查询');link(p,'国家企业信用信息公示系统 https://www.gsxt.gov.cn/', 'https://www.gsxt.gov.cn/')
    p=paragraph(label='人社咨询与办事');link(p,'全国人力资源和社会保障政务服务平台 https://www.12333.gov.cn/', 'https://www.12333.gov.cn/')
    paragraph('人社咨询电话：12333。具体业务以所在地人社部门受理范围为准。',color='71847D')
    document.add_heading('04  依据与来源',1)
    paragraph('法条、风险提示、指南及案例分别标注；案例结论只用于个案参考。',color='71847D')
    kinds={'law':'法律条文','warning':'官方风险提示','guide':'办事与证据指南','case':'案例参考','support':'辅助资料（团队整理）'}
    linked_urls={};groups={};shown_metadata={}
    used={sid for f in findings for sid in f.get('evidence_ids',[])}
    for f in findings:
        for claim in f.get('legal_claims',[]):used.update(claim.get('basis_ids',[]))
    displayed_sources=[s for s in ordered_sources(report) if not findings or s['id'] in used]
    for source in displayed_sources:
        key=(source.get('url') or source['id'],source.get('kind'),source.get('source_publisher'),source.get('source_relation'))
        groups.setdefault(key,[]).append(source)
    for entries in groups.values():
        start=len(document.paragraphs);source=entries[0]
        document.add_heading(f'{source["id"]}  {source["title"]}',2)
        paragraph(kinds.get(source.get('kind'),'资料'),color='71847D')
        details=[]
        for label,value in ((source.get('source_role','发布机构'),source.get('source_publisher')),
                            ('原文载体',source.get('source_site')),
                            ('核对说明',source.get('source_access_note') or source.get('verification_note'))):
            if not value:continue
            key=(label,value)
            if key in shown_metadata:
                details.append(label+'：同 '+shown_metadata[key])
            else:
                details.append(label+'：'+value)
                shown_metadata[key]=source['id']
        if details:paragraph('；'.join(details),color='71847D')
        dates=sorted({e['verified_at'] for e in entries if e.get('verified_at')})
        if dates:paragraph('、'.join(dates),label='资料核对日期',color='71847D')
        # Shared links point to the explicitly named earlier source.
        if source.get('url'):
            url=source['url']
            if url not in linked_urls:
                label='关联官方依据' if source.get('source_relation')=='related_official_basis' else '原文链接'
                p=paragraph(label=label);link(p,'点击打开官方原文',url)
                linked_urls[url]=source['id']
                for alternate in source.get('alternate_urls',[]):
                    p=paragraph(label='备用官方页面');link(p,'点击打开备用官方页面',alternate)
            else:paragraph('原文链接：同 '+linked_urls[url]+'。',color='71847D')
        elif source.get('source_document'):paragraph(source['source_document']+' · '+source.get('source_paragraphs',''),label='本地来源')
        if source.get('legal_database_url') and source.get('legal_database_status')!='specific_record_checked':
            p=paragraph(label='法律数据库查询入口');link(p,'国家法律法规数据库（查询入口，未核验具体记录）',source['legal_database_url'])
        boundaries={}
        for entry in entries:
            name=entry['title']+' · ' if entry['title']!=source['title'] else ''
            p=paragraph(f'{entry["id"]} · {name}{entry["location"]}')
            startmark=OxmlElement('w:bookmarkStart');startmark.set(qn('w:id'),str(1000+displayed_sources.index(entry)));startmark.set(qn('w:name'),'source_'+entry['id'].replace('-','_'))
            endmark=OxmlElement('w:bookmarkEnd');endmark.set(qn('w:id'),startmark.get(qn('w:id')))
            p._p.insert(0,startmark);p._p.append(endmark)
            for run in p.runs:run.bold=True
            p.paragraph_format.keep_with_next=True
            if entry.get('text'):paragraph(entry['text'])
            if entry.get('boundary'):boundaries.setdefault(entry['boundary'],[]).append(entry['id'])
            if entry.get('verification_status')=='local_document_checked_official_pending':paragraph('仅对照用户提供的本地文书，官方真实性与完整性待核验。',label='核验状态',color='B78224')
        for boundary,ids in boundaries.items():paragraph(boundary,label='适用边界 '+ '、'.join(ids))
        for p in document.paragraphs[start+1:]:
            p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.size=Pt(9.5)
        document.paragraphs[start].paragraph_format.keep_with_next=True
    buffer=BytesIO(); document.save(buffer); return buffer.getvalue()
